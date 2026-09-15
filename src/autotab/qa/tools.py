"""Read-only, workbook-scoped QA tools."""

from __future__ import annotations

import re
from collections.abc import Callable
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


class WorkbookToolError(ValueError):
    """Raised for invalid, safely reportable workbook tool input."""


_A1_RANGE = re.compile(r"\A[A-Za-z]{1,3}[1-9]\d*(?::[A-Za-z]{1,3}[1-9]\d*)?\Z")


class WorkbookSession:
    """Own one read-only workbook and expose its approved QA tools."""

    def __init__(self, workbook_path: str | Path, max_range_cells: int) -> None:
        path = Path(workbook_path)
        if not path.is_file():
            raise WorkbookToolError("Workbook does not exist")
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise WorkbookToolError("Workbook must be an .xlsx or .xlsm file")
        if not isinstance(max_range_cells, int) or isinstance(max_range_cells, bool):
            raise WorkbookToolError("max_range_cells must be a positive integer")
        if max_range_cells <= 0:
            raise WorkbookToolError("max_range_cells must be a positive integer")
        self._workbook: Workbook = load_workbook(path, read_only=False, data_only=True)
        self._max_range_cells = max_range_cells

    @property
    def tools(self) -> dict[str, Callable[..., object]]:
        """Return the fixed tool registry for this session."""
        return {
            "load_dataframe": self.load_dataframe,
            "inspect_range": self.inspect_range,
        }

    @property
    def description(self) -> str:
        """Return safe worksheet names and dimensions for the QA prompt."""
        return ", ".join(
            f"{sheet.title} ({rows} rows x {columns} columns)"
            for sheet in self._workbook.worksheets
            for rows, columns in [_worksheet_dimensions(sheet)]
        )

    def close(self) -> None:
        """Close the workbook without saving it."""
        self._workbook.close()

    def load_dataframe(
        self,
        sheet_name: str | None = None,
        has_headers: bool = False,
        range_ref: str | None = None,
    ) -> pd.DataFrame:
        """Load a worksheet table as a DataFrame.

        Args:
            sheet_name: Exact worksheet name, or None for the active worksheet.
            has_headers: Whether to use the table's header rows as column labels.
            range_ref: Optional A1 range bounding the table, or None for the full worksheet.

        Returns:
            Worksheet values as a DataFrame.

        Raises:
            WorkbookToolError: If the worksheet name or range is invalid.
        """
        sheet = self._sheet(sheet_name)
        min_col, min_row, max_col, max_row = _table_bounds(sheet, range_ref)
        if max_row == 0 or max_col == 0:
            return pd.DataFrame()
        if has_headers:
            rows = _expanded_rows(sheet, min_row, max_row, min_col, max_col)
        else:
            rows = [
                list(row)
                for row in sheet.iter_rows(
                    min_row=min_row,
                    max_row=max_row,
                    min_col=min_col,
                    max_col=max_col,
                    values_only=True,
                )
            ]
        if not rows:
            return pd.DataFrame()
        if has_headers:
            header_rows = _header_row_count(sheet, min_row, max_row, min_col, max_col)
            if header_rows >= len(rows):
                return pd.DataFrame()
            headers = _column_headers(rows[:header_rows])
            return pd.DataFrame(rows[header_rows:], columns=headers)
        return pd.DataFrame(rows)

    def inspect_range(
        self, range_ref: str, sheet_name: str | None = None
    ) -> list[list[object | None]]:
        """Read values from an A1 range on the active or named worksheet.

        Args:
            range_ref: A single cell or rectangular A1 range.
            sheet_name: Exact worksheet name, or None for the active worksheet.

        Returns:
            Cell values as a two-dimensional row-major list.

        Raises:
            WorkbookToolError: If the sheet, range, bounds, or range size is invalid.
        """
        if not isinstance(range_ref, str) or not _A1_RANGE.fullmatch(range_ref):
            raise WorkbookToolError("range_ref must be a valid A1 range")
        sheet = self._sheet(sheet_name)
        min_col, min_row, max_col, max_row = range_boundaries(range_ref.upper())
        if min_col > max_col or min_row > max_row:
            raise WorkbookToolError("range_ref must start at its top-left cell")
        sheet_rows, sheet_columns = _worksheet_dimensions(sheet)
        if max_row > sheet_rows or max_col > sheet_columns:
            raise WorkbookToolError("range_ref is outside the worksheet bounds")
        cell_count = (max_row - min_row + 1) * (max_col - min_col + 1)
        if cell_count > self._max_range_cells:
            raise WorkbookToolError(f"range_ref exceeds the {self._max_range_cells}-cell limit")
        return [
            list(row)
            for row in sheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
                values_only=True,
            )
        ]

    def _sheet(self, sheet_name: str | None) -> Worksheet:
        if sheet_name is None:
            return self._workbook.active
        if not isinstance(sheet_name, str) or sheet_name not in self._workbook.sheetnames:
            raise WorkbookToolError(f"Unknown worksheet: {sheet_name!r}")
        return self._workbook[sheet_name]


def _table_bounds(sheet: Worksheet, range_ref: str | None) -> tuple[int, int, int, int]:
    sheet_rows, sheet_columns = _worksheet_dimensions(sheet)
    if range_ref is None:
        return 1, 1, sheet_columns, sheet_rows
    if not isinstance(range_ref, str) or not _A1_RANGE.fullmatch(range_ref):
        raise WorkbookToolError("range_ref must be a valid A1 range")
    min_col, min_row, max_col, max_row = range_boundaries(range_ref.upper())
    if max_row > sheet_rows or max_col > sheet_columns:
        raise WorkbookToolError("range_ref is outside the worksheet bounds")
    return min_col, min_row, max_col, max_row


def _expanded_rows(
    sheet: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int
) -> list[list[object | None]]:
    rows = [
        list(row)
        for row in sheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
            values_only=True,
        )
    ]
    for merged_range in sheet.merged_cells.ranges:
        if (
            merged_range.max_row < min_row
            or merged_range.min_row > max_row
            or merged_range.max_col < min_col
            or merged_range.min_col > max_col
        ):
            continue
        value = sheet.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(
            max(min_row, merged_range.min_row), min(max_row, merged_range.max_row) + 1
        ):
            for column in range(
                max(min_col, merged_range.min_col), min(max_col, merged_range.max_col) + 1
            ):
                rows[row - min_row][column - min_col] = value
    return rows


def _header_row_count(
    sheet: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int
) -> int:
    header_end = min_row
    for merged_range in sheet.merged_cells.ranges:
        if (
            merged_range.min_row == min_row
            and merged_range.max_col >= min_col
            and merged_range.min_col <= max_col
        ):
            header_end = max(header_end, min(max_row, merged_range.max_row))
    return header_end - min_row + 1


def _column_headers(header_rows: list[list[object | None]]) -> list[str]:
    paths = [
        [str(row[column]).strip() for row in header_rows if row[column] not in (None, "")]
        for column in range(len(header_rows[0]))
    ]
    headers = [path[-1] if path else f"Col{index}" for index, path in enumerate(paths)]
    duplicates = {header for header in headers if headers.count(header) > 1}
    for index, header in enumerate(headers):
        if header in duplicates:
            headers[index] = " / ".join(dict.fromkeys(paths[index])) or f"Col{index}"
    for index, header in enumerate(headers):
        duplicate_index = headers[:index].count(header)
        if duplicate_index:
            headers[index] = f"{header} ({duplicate_index + 1})"
    return headers


def _worksheet_dimensions(sheet) -> tuple[int, int]:
    rows, columns = sheet.max_row or 0, sheet.max_column or 0
    if rows == columns == 1 and sheet["A1"].value is None:
        return 0, 0
    return rows, columns
