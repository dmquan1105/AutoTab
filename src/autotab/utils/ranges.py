"""A1 range parsing and viewport clipping."""

from __future__ import annotations

from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter


def parse_cell(cell: str) -> tuple[int, int]:
    """Return row and column for an A1 cell."""
    col, row = coordinate_from_string(cell.upper())
    return int(row), column_index_from_string(col)


def viewport(anchor: str, rows: int, cols: int, size: int) -> str:
    """Create one clipped odd-sized viewport within worksheet dimensions."""
    if size <= 0 or size % 2 == 0:
        raise ValueError("size must be positive and odd")
    row, col = parse_cell(anchor)
    half = size // 2
    r1, c1 = max(1, row - half), max(1, col - half)
    r2, c2 = min(rows, row + half), min(cols, col + half)
    return f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
