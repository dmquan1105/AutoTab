"""Spreadsheet viewport rendering utilities."""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
import threading
import unicodedata
from copy import copy
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from PIL import Image, ImageChops, ImageDraw

from .pdfium_worker import bounded_scale


class WindowRenderer:
    """Render a workbook viewport while preserving workbook formatting."""

    def __init__(
        self,
        workers: int = 1,
        *,
        libreoffice_path: str | Path | None = None,
        timeout_seconds: float = 60,
        image_resolution: int = 300,
        max_image_dimension: int = 8192,
        max_image_pixels: int = 33_554_432,
    ) -> None:
        if not isinstance(workers, int) or isinstance(workers, bool) or workers <= 0:
            raise ValueError("workers must be a positive integer")
        if (
            not isinstance(timeout_seconds, (int, float))
            or isinstance(timeout_seconds, bool)
            or timeout_seconds <= 0
        ):
            raise ValueError("timeout_seconds must be positive")
        for name, value in (
            ("image_resolution", image_resolution),
            ("max_image_dimension", max_image_dimension),
            ("max_image_pixels", max_image_pixels),
        ):
            if not isinstance(value, int) or isinstance(value, bool) or value <= 0:
                raise ValueError(f"{name} must be a positive integer")
        self.workers = workers
        self.libreoffice_path = libreoffice_path
        self.timeout_seconds = float(timeout_seconds)
        self.image_resolution = image_resolution
        self.max_image_dimension = max_image_dimension
        self.max_image_pixels = max_image_pixels
        self._semaphore = threading.BoundedSemaphore(workers)

    def render(
        self, workbook: str | Path, sheet: str, cell_range: str, output: Path, metadata: dict
    ) -> dict:
        warning = None
        with self._semaphore:
            try:
                renderer = _render_with_libreoffice(
                    workbook,
                    sheet,
                    cell_range,
                    output,
                    max_workers=self.workers,
                    libreoffice_path=self.libreoffice_path,
                    timeout_seconds=self.timeout_seconds,
                    image_resolution=self.image_resolution,
                    max_image_dimension=self.max_image_dimension,
                    max_image_pixels=self.max_image_pixels,
                )
            except subprocess.TimeoutExpired as exc:
                warning = f"Viewport rendering timed out after {exc.timeout} seconds"
                renderer = _render_with_pillow(workbook, sheet, cell_range, output)
            except (FileNotFoundError, RuntimeError, ImportError):
                renderer = _render_with_pillow(workbook, sheet, cell_range, output)
            _constrain_image(output, self.max_image_dimension, self.max_image_pixels)
        with Image.open(output) as image:
            width, height = image.size
        metadata.update(
            {
                "sheet": sheet,
                "range": cell_range,
                "image_width": width,
                "image_height": height,
                "renderer": renderer,
                "resolution": self.image_resolution if renderer == "libreoffice" else 1,
                "coordinate_visibility": True,
            }
        )
        if warning:
            metadata["render_warning"] = warning
        return metadata


def _render_with_libreoffice(
    workbook: str | Path,
    sheet: str,
    cell_range: str,
    output: Path,
    *,
    max_workers: int = 1,
    libreoffice_path: str | Path | None = None,
    timeout_seconds: float = 60,
    image_resolution: int = 300,
    max_image_dimension: int = 8192,
    max_image_pixels: int = 33_554_432,
) -> str:
    soffice = str(libreoffice_path) if libreoffice_path else None
    if soffice and not Path(soffice).is_file():
        raise FileNotFoundError(f"LibreOffice executable not found: {soffice}")
    # On Windows soffice.exe detaches and can return before the PDF exists, so the
    # console variant soffice.com is preferred: it waits and reports a real exit code.
    if not soffice and sys.platform == "win32":
        soffice = shutil.which("soffice.com")
    soffice = soffice or shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        for candidate in (
            r"C:\Program Files\LibreOffice\program\soffice.com",
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.com",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ):
            if Path(candidate).is_file():
                soffice = candidate
                break
    if not soffice:
        raise FileNotFoundError("LibreOffice executable not found")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="autotab_render_") as temp:
        temp_dir = Path(temp)
        prepared = temp_dir / Path(workbook).name
        pdf = prepared.with_suffix(".pdf")
        profile = temp_dir / "profile"
        wb = openpyxl.load_workbook(workbook)
        try:
            ws = wb[sheet]
            wb.active = wb.worksheets.index(ws)
            for other in wb.worksheets:
                other.sheet_state = "visible" if other is ws else "hidden"
            ws.print_area = cell_range
            ws.print_options.headings = True
            ws.print_options.gridLines = True
            ws.page_margins.left = ws.page_margins.right = ws.page_margins.top = (
                ws.page_margins.bottom
            ) = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
            _prepare_dimensions(ws, cell_range)
            wb.save(prepared)
        finally:
            wb.close()
        command = [
            str(soffice),
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--nodefault",
            "--nolockcheck",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--convert-to",
            "pdf:calc_pdf_Export",
            "--outdir",
            str(temp_dir),
            str(prepared),
        ]
        result = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=max(1, float(timeout_seconds)),
            check=False,
        )
        if result.returncode or not pdf.is_file():
            raise RuntimeError(
                result.stderr.strip() or result.stdout.strip() or "LibreOffice failed"
            )
        if max_workers > 1:
            _render_pdf_page_in_subprocess(
                pdf,
                output,
                image_resolution,
                timeout_seconds=timeout_seconds,
                max_image_dimension=max_image_dimension,
                max_image_pixels=max_image_pixels,
            )
        else:
            # PDFium APIs are process-global and not thread-safe.
            with _PDFIUM_LOCK:
                try:
                    import pypdfium2
                except ImportError as exc:
                    raise ImportError("pypdfium2 is required for LibreOffice rendering") from exc
                document = pypdfium2.PdfDocument(str(pdf))
                try:
                    page = document[0]
                    try:
                        width, height = page.get_size()
                        scale = bounded_scale(
                            width,
                            height,
                            image_resolution,
                            max_image_dimension,
                            max_image_pixels,
                        )
                        bitmap = page.render(scale=scale)
                        try:
                            bitmap.to_pil().save(output)
                        finally:
                            bitmap.close()
                    finally:
                        page.close()
                finally:
                    document.close()
    _trim_image(output)
    return "libreoffice"


_PDFIUM_LOCK = threading.Lock()


def _render_pdf_page_in_subprocess(
    pdf_path: Path,
    output: Path,
    resolution: int,
    *,
    timeout_seconds: float,
    max_image_dimension: int,
    max_image_pixels: int,
) -> None:
    worker_path = Path(__file__).with_name("pdfium_worker.py").resolve()
    result = subprocess.run(
        [
            sys.executable,
            str(worker_path),
            str(pdf_path),
            str(output),
            str(resolution),
            str(max_image_dimension),
            str(max_image_pixels),
        ],
        capture_output=True,
        text=True,
        timeout=max(1, float(timeout_seconds)),
        check=False,
    )
    if result.returncode or not output.is_file():
        detail = result.stderr.strip() or result.stdout.strip()
        raise RuntimeError("PDFium worker failed" + (f": {detail}" if detail else ""))


def _prepare_dimensions(ws: Any, cell_range: str) -> None:
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    for row in range(min_row, max_row + 1):
        ws.row_dimensions[row].height = max(float(ws.row_dimensions[row].height or 15), 18)
    for col in range(min_col, max_col + 1):
        dimension = ws.column_dimensions[get_column_letter(col)]
        width = float(dimension.width or 13)
        for row in range(min_row, max_row + 1):
            value = ws.cell(row, col).value
            if value is not None:
                width = max(width, min(50, _display_width(str(value)) + 3))
        dimension.width = width
    merged_by_cell = {}
    for merged in ws.merged_cells.ranges:
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                merged_by_cell[(row, col)] = merged
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            merged = merged_by_cell.get((cell.row, cell.column))
            columns = (
                range(merged.min_col, merged.max_col + 1)
                if merged
                else range(cell.column, cell.column + 1)
            )
            available: float = sum(
                float(ws.column_dimensions[get_column_letter(col)].width or 13) for col in columns
            )
            if cell.value is not None and _display_width(str(cell.value)) > available:
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                cell.alignment = alignment
            if cell.value is not None and (cell.alignment.wrap_text or "\n" in str(cell.value)):
                lines = str(cell.value).splitlines() or [""]
                line_count = sum(
                    max(1, (len(line) + max(1, int(available)) - 1) // max(1, int(available)))
                    for line in lines
                )
                ws.row_dimensions[cell.row].height = max(
                    float(ws.row_dimensions[cell.row].height or 18),
                    line_count * float(cell.font.sz or 11) * 1.25 + 4,
                )


def _render_with_pillow(workbook: str | Path, sheet: str, cell_range: str, output: Path) -> str:
    wb = openpyxl.load_workbook(workbook, data_only=False)
    ws = wb[sheet]
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    values = [
        [ws.cell(r, c).value for c in range(min_col, max_col + 1)]
        for r in range(min_row, max_row + 1)
    ]
    widths = [
        max(90, min(360, max([_display_width(str(row[i] or "")) for row in values] + [8]) * 8 + 24))
        for i in range(max_col - min_col + 1)
    ]
    row_height = 30
    image = Image.new("RGB", (sum(widths) + 2, max(40, len(values) * row_height + 2)), "white")
    draw = ImageDraw.Draw(image)
    y = 1
    for row_index, row in enumerate(values, min_row):
        x = 1
        for i, value in enumerate(row):
            cell = ws.cell(row_index, min_col + i)
            fill = (
                cell.fill.fgColor.rgb
                if cell.fill.fill_type and cell.fill.fgColor.type == "rgb"
                else None
            )
            color = (
                f"#{fill[-6:]}"
                if fill and len(fill) >= 6
                else ("#eaf2f8" if row_index == min_row else "white")
            )
            draw.rectangle((x, y, x + widths[i], y + row_height), fill=color, outline="#b7c3d0")
            draw.text(
                (x + 8, y + 7), str(value or "")[: max(1, widths[i] // 8 - 2)], fill="#17202a"
            )
            x += widths[i]
        y += row_height
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output)
    return "pillow_fallback"


def _trim_image(path: Path) -> None:
    with Image.open(path) as image:
        rgb = image.convert("RGB")
        bbox = ImageChops.difference(rgb, Image.new("RGB", rgb.size, "white")).getbbox()
        if bbox:
            rgb.crop(
                (
                    max(0, bbox[0] - 6),
                    max(0, bbox[1] - 6),
                    min(rgb.width, bbox[2] + 6),
                    min(rgb.height, bbox[3] + 6),
                )
            ).save(path)


def _constrain_image(path: Path, max_dimension: int, max_pixels: int) -> None:
    with Image.open(path) as image:
        width, height = image.size
        factor = min(
            1.0,
            max_dimension / max(1, width),
            max_dimension / max(1, height),
            (max_pixels / max(1, width * height)) ** 0.5,
        )
        if factor < 1:
            size = (max(1, int(width * factor)), max(1, int(height * factor)))
            image.resize(size, Image.Resampling.LANCZOS).save(path)


def _display_width(value: str) -> int:
    return sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in value)
