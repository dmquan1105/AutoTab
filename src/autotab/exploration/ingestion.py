"""Workbook snapshot ingestion."""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

import openpyxl


@dataclass(frozen=True)
class CellDocument:
    workbook: str
    sheet: str
    coordinate: str
    text: str
    raw_value: Any
    displayed_value: str
    formula: str | None
    number_format: str
    headers: tuple[str, ...] = ()
    table_metadata: dict[str, Any] | None = None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class WorkbookSnapshot:
    path: str
    sheets: tuple[dict[str, Any], ...]
    cells: tuple[CellDocument, ...]


class WorkbookLoader:
    """Load workbook structure and searchable cell documents."""

    def load(self, path: str | Path, include_hidden: bool = False) -> WorkbookSnapshot:
        source = Path(path)
        wb_formula = openpyxl.load_workbook(source, data_only=False)
        wb_values = openpyxl.load_workbook(source, data_only=True)
        sheets: list[dict[str, Any]] = []
        docs: list[CellDocument] = []
        for ws in wb_formula.worksheets:
            if ws.sheet_state != "visible" and not include_hidden:
                continue
            value_ws = wb_values[ws.title]
            sheets.append(
                {
                    "name": ws.title,
                    "state": ws.sheet_state,
                    "max_row": ws.max_row,
                    "max_column": ws.max_column,
                    "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
                }
            )
            for row in ws.iter_rows():
                for cell in row:
                    raw = cell.value
                    shown = value_ws[cell.coordinate].value
                    if raw is None and shown is None:
                        continue
                    formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                    display = "" if shown is None else str(shown)
                    docs.append(
                        CellDocument(
                            source.name,
                            ws.title,
                            cell.coordinate,
                            display,
                            raw,
                            display,
                            formula,
                            cell.number_format,
                        )
                    )
        return WorkbookSnapshot(str(source), tuple(sheets), tuple(docs))


def write_snapshot(snapshot: WorkbookSnapshot, path: Path) -> None:
    """Write a JSON debugging snapshot."""
    path.write_text(
        json.dumps(
            {
                "path": snapshot.path,
                "sheets": snapshot.sheets,
                "cells": [c.to_dict() for c in snapshot.cells],
            },
            indent=2,
        ),
        encoding="utf-8",
    )
