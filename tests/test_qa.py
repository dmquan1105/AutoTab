from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from autotab import cli
from autotab.config import ConfigError, load_config
from autotab.qa import worker
from autotab.qa.agent import QAAgent
from autotab.qa.parser import ResponseFormatError, parse_response
from autotab.qa.sandbox import Sandbox
from autotab.qa.tools import WorkbookSession, WorkbookToolError
from autotab.qa.types import Action, FinalAnswer


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    """Create a small multi-sheet workbook for QA tests."""
    path = tmp_path / "qa.xlsx"
    workbook = Workbook()
    active = workbook.active
    active.title = "Sales"
    active.append(["Region", "Amount", "Note"])
    active.append(["North", 10, None])
    active.append(["South", 20, "ok"])
    other = workbook.create_sheet("Other")
    other.append(["Name", "Value"])
    other.append(["A", 3])
    workbook.create_sheet("Empty")
    workbook.save(path)
    workbook.close()
    return path


def test_qa_config_defaults_and_validation(tmp_path: Path) -> None:
    config = load_config(tmp_path / "missing.yaml")
    assert config["qa"]["exploration_enabled"] is True
    assert config["qa"]["max_turns"] == 10
    invalid = tmp_path / "invalid.yaml"
    invalid.write_text("qa:\n  extra: true\n", encoding="utf-8")
    with pytest.raises(ConfigError, match="Unknown qa key"):
        load_config(invalid)
    invalid.write_text("qa:\n  max_turns: false\n", encoding="utf-8")
    with pytest.raises(ConfigError, match="positive integer"):
        load_config(invalid)
    invalid.write_text("qa:\n  exploration_enabled: enabled\n", encoding="utf-8")
    with pytest.raises(ConfigError, match="must be boolean"):
        load_config(invalid)


def test_workbook_tools(workbook_path: Path) -> None:
    session = WorkbookSession(workbook_path, max_range_cells=6)
    try:
        frame = session.load_dataframe()
        assert list(frame.columns) == ["Region", "Amount", "Note"]
        assert frame.loc[0, "Region"] == "North"
        assert pd.isna(frame.loc[0, "Note"])
        assert session.load_dataframe("Other").to_dict("records") == [{"Name": "A", "Value": 3}]
        assert session.inspect_range("B2") == [[10]]
        assert session.inspect_range("A1:B2", "Sales") == [
            ["Region", "Amount"],
            ["North", 10],
        ]
        assert "Empty (0 rows x 0 columns)" in session.description
        assert set(session.tools) == {"load_dataframe", "inspect_range"}
    finally:
        session.close()


@pytest.mark.parametrize("range_ref", ["", "A0", "A", "A1:", "B2:A1"])
def test_inspect_range_rejects_invalid_ranges(workbook_path: Path, range_ref: str) -> None:
    session = WorkbookSession(workbook_path, max_range_cells=3)
    try:
        with pytest.raises(WorkbookToolError):
            session.inspect_range(range_ref)
    finally:
        session.close()


def test_workbook_tools_reject_sheet_bounds_and_size(workbook_path: Path) -> None:
    session = WorkbookSession(workbook_path, max_range_cells=3)
    try:
        with pytest.raises(WorkbookToolError, match="Unknown worksheet"):
            session.load_dataframe("Missing")
        with pytest.raises(WorkbookToolError, match="outside"):
            session.inspect_range("D1")
        with pytest.raises(WorkbookToolError, match="cell limit"):
            session.inspect_range("A1:B2")
        with pytest.raises(WorkbookToolError, match="outside"):
            session.inspect_range("A1", "Empty")
    finally:
        session.close()


def test_parser_accepts_one_strict_response_form() -> None:
    action = parse_response("Thought: inspect\nAction: ```python\nprint(inspect_range('A1'))\n```")
    final = parse_response("Thought: grounded\nFinal Answer: 10")
    assert action == Action("inspect", "print(inspect_range('A1'))")
    assert final == FinalAnswer("grounded", "10")


@pytest.mark.parametrize(
    "response",
    [
        "Final Answer: 10",
        "Thought: both\nAction: ```python\nprint(1)\n```\nFinal Answer: 1",
        "Thought: two\nAction: ```python\nprint(1)\n```\n```python\nprint(2)\n```",
        "Thought: empty\nAction: ```python\n\n```",
    ],
)
def test_parser_rejects_invalid_responses(response: str) -> None:
    with pytest.raises(ResponseFormatError):
        parse_response(response)


def test_sandbox_allows_tools_and_dataframe_operations(workbook_path: Path) -> None:
    before = workbook_path.read_bytes()
    sandbox = Sandbox(workbook_path, 10, 512, 2000, 100)
    result = sandbox.execute(
        "df = load_dataframe('Sales')\n"
        "ng = df[df['Amount'] > 10]\n"
        "print(ng[['Region', 'Amount']].to_dict('records'))"
    )
    assert result.success
    assert "South" in result.text
    assert workbook_path.read_bytes() == before


def test_sandbox_captures_unicode_output(workbook_path: Path) -> None:
    value = "Ti\u1ebfng Vi\u1ec7t"
    result = Sandbox(workbook_path, 10, 512, 2000, 100).execute(f"print({value!r})")
    assert result.success
    assert result.text == value


@pytest.mark.parametrize(
    "code, message",
    [
        ("import os", "Import is not allowed"),
        ("print(open('secret.txt'))", "Call to 'open' is not allowed"),
        ("print(__import__('socket'))", "Call to '__import__' is not allowed"),
        ("print(load_dataframe().to_csv())", "Attribute 'to_csv' is not allowed"),
        ("print(eval('1 + 1'))", "Call to 'eval' is not allowed"),
    ],
)
def test_sandbox_rejects_unsafe_code(workbook_path: Path, code: str, message: str) -> None:
    before = workbook_path.read_bytes()
    result = Sandbox(workbook_path, 10, 512, 2000, 100).execute(code)
    assert not result.success
    assert message in result.text
    assert workbook_path.read_bytes() == before


def test_agent_observes_action_then_returns_answer(workbook_path: Path, tmp_path: Path) -> None:
    class FakeModel:
        def __init__(self) -> None:
            self.prompts: list[str] = []

        def complete(self, prompt: str, image_path: str | Path | None = None) -> str:
            self.prompts.append(prompt)
            if len(self.prompts) == 1:
                return (
                    "Thought: inspect the amount\n"
                    "Action: ```python\nprint(inspect_range('B2'))\n```"
                )
            assert "Observation: [[10]]" in prompt
            return "Thought: the observed value answers it\nFinal Answer: 10"

    model = FakeModel()
    result = QAAgent(model, load_config("config.example.yaml")).run(
        workbook_path,
        "What is the first amount?",
        "Exploration found Sales!B2.",
        tmp_path / "qa",
    )
    assert result.success
    assert result.answer == "10"
    assert result.turns == 2
    first_turn = tmp_path / "qa" / "turn-01"
    assert "Exploration found Sales!B2." in (first_turn / "prompt.txt").read_text(encoding="utf-8")
    assert (first_turn / "response.txt").is_file()
    assert (first_turn / "observation.txt").read_text(encoding="utf-8") == "Observation: [[10]]"
    assert json.loads((tmp_path / "qa" / "result.json").read_text(encoding="utf-8")) == {
        "success": True,
        "answer": "10",
        "turns": 2,
    }


def test_agent_recovers_from_format_error_and_bounds_turns(workbook_path: Path) -> None:
    class FakeModel:
        def complete(self, prompt: str, image_path: str | Path | None = None) -> str:
            return "invalid"

    config = load_config("config.example.yaml")
    config["qa"]["max_turns"] = 2
    result = QAAgent(FakeModel(), config).run(workbook_path, "Question")
    assert not result.success
    assert result.turns == 2


def test_cli_runs_qa_workbook_and_query(workbook_path: Path, monkeypatch, capsys) -> None:
    run_root = workbook_path.parent / "run"

    class FakeExplorationPipeline:
        def __init__(self, config: dict) -> None:
            self.config = config

        def run(self, workbooks: list[str], query: str) -> Path:
            path = run_root / "exploration" / "exploration.md"
            path.parent.mkdir(parents=True)
            path.write_text("Grounded exploration evidence", encoding="utf-8")
            return path

    class FakeClient:
        def __init__(self, *args: object, **kwargs: object) -> None:
            self.turn = 0

        def complete(self, prompt: str, image_path: str | Path | None = None) -> str:
            self.turn += 1
            if self.turn == 1:
                return "Thought: inspect\nAction: ```python\nprint(inspect_range('B2'))\n```"
            return "Thought: grounded\nFinal Answer: 10"

    monkeypatch.setattr(cli, "OpenAICompatibleClient", FakeClient)
    monkeypatch.setattr(cli, "ExplorationPipeline", FakeExplorationPipeline)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "autotab",
            "qa",
            "--config",
            "config.example.yaml",
            "--workbook",
            str(workbook_path),
            "--query",
            "What is the first amount?",
        ],
    )
    assert cli.main() == 0
    assert capsys.readouterr().out.strip() == "10"
    assert "Grounded exploration evidence" in (
        run_root / "qa" / "turn-01" / "prompt.txt"
    ).read_text(encoding="utf-8")


def test_cli_can_disable_exploration(workbook_path: Path, tmp_path: Path, monkeypatch) -> None:
    class DisabledExplorationPipeline:
        def __init__(self, config: dict) -> None:
            raise AssertionError("Exploration must remain disabled")

    class FakeClient:
        def __init__(self, *args: object, **kwargs: object) -> None:
            pass

        def complete(self, prompt: str, image_path: str | Path | None = None) -> str:
            return "Thought: done\nFinal Answer: answer"

    config_path = tmp_path / "config.yaml"
    artifact_root = (tmp_path / "outputs").as_posix()
    config_path.write_text(
        "models:\n"
        "  llm:\n"
        "    model: fake\n"
        "    base_url: http://localhost/v1\n"
        f"runtime:\n  artifact_root: {artifact_root}\n"
        "qa:\n  exploration_enabled: false\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(cli, "OpenAICompatibleClient", FakeClient)
    monkeypatch.setattr(cli, "ExplorationPipeline", DisabledExplorationPipeline)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "autotab",
            "qa",
            "--config",
            str(config_path),
            "--workbook",
            str(workbook_path),
            "--query",
            "Question",
        ],
    )

    assert cli.main() == 0
    assert len(list((tmp_path / "outputs").glob("*/qa/result.json"))) == 1


def test_worker_executes_with_bounded_output(workbook_path: Path) -> None:
    payload = {
        "workbook_path": str(workbook_path),
        "max_observation_chars": 5,
        "max_range_cells": 100,
        "code": "print(inspect_range('A1:B2'))",
    }
    assert worker.execute(payload) == "[['Re"
    writer = worker.LimitedWriter(3)
    assert writer.write("abcdef") == 6
    assert writer.value() == "abc"


def test_worker_main_returns_safe_errors(workbook_path: Path, monkeypatch, capsys) -> None:
    payload = {
        "workbook_path": str(workbook_path),
        "max_observation_chars": 100,
        "max_range_cells": 100,
        "code": "import os",
    }
    monkeypatch.setattr(sys, "stdin", io.StringIO(json.dumps(payload)))
    assert worker.main() == 1
    assert "Import is not allowed" in capsys.readouterr().err


def test_all_sample_query_workbook_cases_complete() -> None:
    cases = [
        ("multitab.xlsx", "query_simple.txt"),
        ("multitab.xlsx", "query_intermediate.txt"),
        ("multitab.xlsx", "query_complex.txt"),
        ("injuries-table14.xlsx", "query_injuries_simple.txt"),
        ("injuries-table14.xlsx", "query_injuries_intermediate.txt"),
        ("injuries-table14.xlsx", "query_injuries_complex.txt"),
    ]

    class FakeModel:
        def __init__(self) -> None:
            self.turn = 0

        def complete(self, prompt: str, image_path: str | Path | None = None) -> str:
            self.turn += 1
            if self.turn == 1:
                return "Thought: inspect\nAction: ```python\nprint(inspect_range('A1'))\n```"
            assert "Observation:" in prompt
            return "Thought: observed workbook data\nFinal Answer: completed"

    config = load_config("config.example.yaml")
    for workbook_name, query_name in cases:
        query = (Path("samples") / query_name).read_text(encoding="utf-8")
        result = QAAgent(FakeModel(), config).run(Path("samples") / workbook_name, query)
        assert result.success, query_name
