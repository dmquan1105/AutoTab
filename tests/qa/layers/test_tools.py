"""Read-only workbook registry behavior; see ``specs/layers/tools.md``."""

import hashlib
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from autotab.qa.layers.tools import (
    AttributeName,
    ToolRegistry,
    WorkbookReader,
)
from autotab.qa.schemas import ResultState, ToolResult

WORKBOOK_ID = "workbook_1"


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    book = Workbook()
    revenue = book.active
    revenue.title = "Revenue"
    revenue["A1"] = "Region"
    revenue["B1"] = "Amount"
    revenue["A2"] = "North"
    revenue["B2"] = 50
    revenue["A3"] = "South"
    revenue["B3"] = 40
    revenue["A4"] = "Total"
    revenue["B4"] = "=SUM(B2:B3)"
    revenue["A6"] = "Q1 Summary"
    revenue.merge_cells("A6:B6")
    revenue["B8"] = 0
    revenue["A9"] = "N/A"
    revenue["B2"].number_format = "#,##0"

    costs = book.create_sheet("Costs")
    costs["A1"] = "Region"
    costs["A2"] = "north"
    costs["A3"] = " North "

    notes = book.create_sheet("Notes")
    notes["A1"] = "internal"
    notes.sheet_state = "hidden"

    path = tmp_path / "book.xlsx"
    book.save(path)
    return path


@pytest.fixture
def registry(workbook_path: Path) -> ToolRegistry:
    reader = WorkbookReader(WORKBOOK_ID, workbook_path, max_range_cells=50)
    return ToolRegistry([reader])


def _dispatch(registry: ToolRegistry, tool: str, **arguments: Any) -> ToolResult:
    return registry.dispatch(tool, {"workbook_id": WORKBOOK_ID, **arguments})


def _cells(result: ToolResult) -> list[dict[str, Any]]:
    assert isinstance(result.result, dict)
    rows = result.result["cells"]
    assert isinstance(rows, list)
    return [cell for row in rows for cell in row]


def test_registry_exposes_exactly_the_v1_tools_with_generated_schemas(
    registry: ToolRegistry,
) -> None:
    descriptions = registry.descriptions()

    assert [item["name"] for item in descriptions] == [
        "inspect_range",
        "inspect_attributes",
        "search",
        "get_sheet_as_dataframe",
    ]
    for item in descriptions:
        assert item["description"].strip()
        assert item["schema"]["properties"]["workbook_id"]
        assert item["schema"]["additionalProperties"] is False


def test_unknown_tool_and_unknown_workbook_are_rejected_without_reading(
    registry: ToolRegistry,
) -> None:
    unknown_tool = registry.dispatch("execute_excel", {"workbook_id": WORKBOOK_ID})
    assert unknown_tool.state is ResultState.ERROR
    assert "execute_excel" in (unknown_tool.error or "")

    unknown_workbook = registry.dispatch(
        "inspect_range", {"workbook_id": "workbook_9", "range_ref": "A1"}
    )
    assert unknown_workbook.state is ResultState.ERROR
    assert "workbook_9" in (unknown_workbook.error or "")


@pytest.mark.parametrize(
    "arguments",
    [
        {},
        {"range_ref": ""},
        {"range_ref": "A1", "sheet_name": "Revenue", "unexpected": 1},
        {"range_ref": "not-a-range"},
        {"range_ref": "A1", "sheet_name": "Missing"},
    ],
)
def test_invalid_arguments_fail_before_execution(
    registry: ToolRegistry, arguments: dict[str, Any]
) -> None:
    result = _dispatch(registry, "inspect_range", **arguments)

    assert result.state is ResultState.ERROR
    assert result.error


def test_inspect_range_reports_resolved_sheet_and_stable_shape(registry: ToolRegistry) -> None:
    single = _dispatch(registry, "inspect_range", range_ref="A2")
    matrix = _dispatch(registry, "inspect_range", range_ref="A1:B3", sheet_name="Revenue")

    assert single.state is ResultState.SUCCESS
    assert isinstance(single.result, dict)
    assert single.result["sheet"] == "Revenue"
    assert single.result["range"] == "A2"
    assert single.result["cells"] == [
        [
            {
                "coordinate": "A2",
                "value": "North",
                "displayed": "North",
                "formula": None,
                "data_type": "s",
                "number_format": "General",
            }
        ]
    ]

    assert isinstance(matrix.result, dict)
    assert len(matrix.result["cells"]) == 3
    assert [cell["coordinate"] for cell in matrix.result["cells"][0]] == ["A1", "B1"]
    assert matrix.provenance[0].sheet == "Revenue"
    assert matrix.provenance[0].range_ref == "A1:B3"


def test_formula_text_stays_separate_and_a_missing_cache_is_not_zero(
    registry: ToolRegistry,
) -> None:
    result = _dispatch(registry, "inspect_range", range_ref="B4")

    cell = _cells(result)[0]

    assert cell["formula"] == "=SUM(B2:B3)"
    assert cell["value"] is None
    assert cell["displayed"] is None


def test_blank_zero_and_not_available_stay_distinct(registry: ToolRegistry) -> None:
    # openpyxl round-trips a written "" to None, so a literal empty string cannot be
    # produced by this fixture; blank, zero, and "N/A" are the distinctions a
    # workbook can actually carry, and none of them may collapse into another.
    values = {
        cell["coordinate"]: cell["value"]
        for cell in _cells(_dispatch(registry, "inspect_range", range_ref="A7:B9"))
    }

    assert values["A7"] is None
    assert values["A8"] is None
    assert values["B8"] == 0
    assert values["A9"] == "N/A"


def test_covered_merged_cells_report_the_anchor_value(registry: ToolRegistry) -> None:
    cells = {
        cell["coordinate"]: cell
        for cell in _cells(_dispatch(registry, "inspect_range", range_ref="A6:B6"))
    }

    assert cells["A6"]["value"] == "Q1 Summary"
    assert cells["B6"]["value"] == "Q1 Summary"
    assert cells["B6"]["merged_with"] == "A6:B6"


def test_hidden_sheets_are_refused_unless_configured(workbook_path: Path) -> None:
    default = ToolRegistry([WorkbookReader(WORKBOOK_ID, workbook_path)])
    permissive = ToolRegistry([WorkbookReader(WORKBOOK_ID, workbook_path, include_hidden=True)])

    refused = _dispatch(default, "inspect_range", range_ref="A1", sheet_name="Notes")
    allowed = _dispatch(permissive, "inspect_range", range_ref="A1", sheet_name="Notes")

    assert refused.state is ResultState.ERROR
    assert "hidden" in (refused.error or "").lower()
    assert allowed.state is ResultState.SUCCESS
    assert "Notes" not in default.sheet_names(WORKBOOK_ID)


def test_oversized_range_is_refused_with_the_limit(registry: ToolRegistry) -> None:
    result = _dispatch(registry, "inspect_range", range_ref="A1:Z100")

    assert result.state is ResultState.ERROR
    assert "50" in (result.error or "")


def test_empty_range_reports_empty_without_inventing_values(registry: ToolRegistry) -> None:
    result = _dispatch(registry, "inspect_range", range_ref="D20:E21")

    assert result.state is ResultState.EMPTY
    assert all(cell["value"] is None for cell in _cells(result))


def test_inspect_attributes_accepts_only_registered_attributes(registry: ToolRegistry) -> None:
    result = _dispatch(
        registry,
        "inspect_attributes",
        range_ref="B2",
        attributes=[AttributeName.NUMBER_FORMAT.value, AttributeName.DATA_TYPE.value],
    )

    assert result.state is ResultState.SUCCESS
    assert _cells(result)[0] == {
        "coordinate": "B2",
        "number_format": "#,##0",
        "data_type": "n",
    }

    for attributes in ([], ["width"], ["formula", "formula"]):
        rejected = _dispatch(registry, "inspect_attributes", range_ref="B2", attributes=attributes)
        assert rejected.state is ResultState.ERROR


@pytest.mark.parametrize(
    ("search_type", "case_sensitive", "expected"),
    [
        ("partial", False, ["Costs!A2", "Costs!A3", "Revenue!A2"]),
        ("partial", True, ["Costs!A3", "Revenue!A2"]),
        ("whole", False, ["Costs!A2", "Revenue!A2"]),
        ("whole", True, ["Revenue!A2"]),
        ("strip", False, ["Costs!A2", "Costs!A3", "Revenue!A2"]),
        ("strip", True, ["Costs!A3", "Revenue!A2"]),
    ],
)
def test_search_matching_modes_are_deterministic(
    registry: ToolRegistry,
    search_type: str,
    case_sensitive: bool,
    expected: list[str],
) -> None:
    result = registry.dispatch(
        "search",
        {
            "workbook_id": WORKBOOK_ID,
            "value": "North",
            "search_type": search_type,
            "case_sensitive": case_sensitive,
        },
    )

    assert isinstance(result.result, dict)
    assert [match["reference"] for match in result.result["matches"]] == expected


def test_search_without_matches_is_empty(registry: ToolRegistry) -> None:
    result = registry.dispatch("search", {"workbook_id": WORKBOOK_ID, "value": "Neverland"})

    assert result.state is ResultState.EMPTY
    assert isinstance(result.result, dict)
    assert result.result["matches"] == []


def test_dataframe_records_selection_and_truncates_explicitly(registry: ToolRegistry) -> None:
    full = _dispatch(registry, "get_sheet_as_dataframe", sheet_name="Revenue")
    capped = _dispatch(registry, "get_sheet_as_dataframe", sheet_name="Revenue", max_rows=2)
    headerless = _dispatch(registry, "get_sheet_as_dataframe", sheet_name="Costs", header_row=None)

    assert isinstance(full.result, dict)
    assert full.result["columns"] == ["Region", "Amount"]
    assert full.result["header_row"] == 1
    assert full.result["rows"][0] == ["North", 50]
    assert full.result["shape"] == [len(full.result["rows"]), 2]

    assert capped.state is ResultState.TRUNCATED
    assert isinstance(capped.result, dict)
    assert len(capped.result["rows"]) == 2

    assert isinstance(headerless.result, dict)
    assert headerless.result["columns"] == ["A", "B"][: headerless.result["shape"][1]]
    assert headerless.result["rows"][0] == ["Region"]


def test_reads_never_modify_the_source_workbook(
    registry: ToolRegistry, workbook_path: Path
) -> None:
    before = hashlib.sha256(workbook_path.read_bytes()).hexdigest()

    _dispatch(registry, "inspect_range", range_ref="A1:B4")
    _dispatch(registry, "get_sheet_as_dataframe", sheet_name="Revenue")
    registry.dispatch("search", {"workbook_id": WORKBOOK_ID, "value": "North"})

    assert hashlib.sha256(workbook_path.read_bytes()).hexdigest() == before


def test_describe_lists_readable_sheets_with_extent_and_merges(workbook_path: Path) -> None:
    reader = WorkbookReader(WORKBOOK_ID, workbook_path)

    sheets = reader.describe()

    assert [sheet["name"] for sheet in sheets] == ["Revenue", "Costs"]
    revenue = sheets[0]
    assert revenue["dimensions"] == "A1:B9"
    assert revenue["merged"] == ["A6:B6"]


# --- structure summary and the tool read limit -----------------------------------------


@pytest.fixture
def scores_path(tmp_path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "Scores"
    sheet["A1"] = "No"
    sheet.merge_cells("A1:A2")
    sheet["B1"] = "Name"
    sheet.merge_cells("B1:C1")
    sheet["B2"] = "First"
    sheet["C2"] = "Last"
    sheet["D1"] = "Score"
    sheet.merge_cells("D1:D2")
    for row, (first, last, score) in enumerate(
        [("Ha", "Dang", 92), ("An", "Pham", 60), ("Minh", None, 99), ("Lan", "Vu", 71)], start=3
    ):
        sheet[f"A{row}"] = row - 2
        sheet[f"B{row}"] = first
        sheet[f"C{row}"] = last
        sheet[f"D{row}"] = score
    sheet["B8"] = "Total"
    sheet["D8"] = "=SUM(D3:D6)"
    path = tmp_path / "scores.xlsx"
    book.save(path)
    return path


def test_describe_profiles_each_column_without_interpreting_the_layout(
    scores_path: Path,
) -> None:
    # Facts about what is where, from every cell; headers and layout are left to the
    # model, since real sheets are too varied for the harness to guess them.
    structure = WorkbookReader(WORKBOOK_ID, scores_path).describe()[0]["structure"]

    assert "header_rows" not in structure and "data_rows" not in structure
    columns = {column["column"]: column for column in structure["columns"]}
    assert "header" not in columns["D"]
    assert columns["D"]["counts"] == {"text": 1, "number": 4}
    assert columns["D"]["numbers"] == {"min": 60, "max": 99, "rows": [3, 6]}
    assert columns["D"]["formulas"] == 1
    assert columns["D"]["text"] == [["D1", "Score"]]
    assert columns["B"]["text"] == [["B1", "Name"], ["B2", "First"], ["B3", "Ha"]]
    assert columns["C"]["counts"] == {"text": 4}
    assert structure["total_rows"] == [8]
    assert structure["blank_rows"] == [7]


def test_describe_caps_the_columns_it_profiles(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    for col in range(1, 71):
        sheet.cell(row=1, column=col, value=f"c{col}")
        sheet.cell(row=2, column=col, value=col)
    path = tmp_path / "wide.xlsx"
    book.save(path)

    structure = WorkbookReader(WORKBOOK_ID, path).describe()[0]["structure"]

    assert len(structure["columns"]) == 40
    assert structure["more_columns"] == 30


def test_a_tool_read_above_the_read_limit_is_refused_with_the_code_route(
    scores_path: Path,
) -> None:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, scores_path)], max_read_cells=10)

    wide = _dispatch(registry, "inspect_range", sheet_name="Scores", range_ref="A1:D8")
    sheet = _dispatch(registry, "get_sheet_as_dataframe", sheet_name="Scores")
    small = _dispatch(registry, "inspect_range", sheet_name="Scores", range_ref="A1:D2")

    assert wide.state is ResultState.ERROR
    assert "32 cells" in (wide.error or "") and "10" in (wide.error or "")
    assert "code" in (wide.error or "") and "wb.sheet('Scores'" in (wide.error or "")
    assert sheet.state is ResultState.ERROR and "code" in (sheet.error or "")
    assert small.state is ResultState.SUCCESS


def test_without_a_read_limit_the_registry_reads_whole_ranges(scores_path: Path) -> None:
    # The sandbox facade uses a registry without the limit: code is where bulk data goes.
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, scores_path)])

    result = _dispatch(registry, "inspect_range", sheet_name="Scores", range_ref="A1:D8")

    assert result.state is ResultState.SUCCESS


def test_tool_descriptions_state_the_read_limit(scores_path: Path) -> None:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, scores_path)], max_read_cells=10)

    described = {item["name"]: item["description"] for item in registry.descriptions()}

    assert "at most 10 cells" in described["inspect_range"]


def test_describe_keeps_the_top_left_cells_verbatim(scores_path: Path) -> None:
    # The shape of a table's top, exactly as the cells hold it; nothing is labelled.
    window = WorkbookReader(WORKBOOK_ID, scores_path).describe()[0]["structure"]["window"]

    assert window["columns"] == ["A", "B", "C", "D"]
    assert window["rows"][0] == [1, ["No", "Name", None, "Score"]]
    assert window["rows"][1] == [2, [None, "First", "Last", None]]
    assert window["rows"][2] == [3, [1, "Ha", "Dang", 92]]
    # Row 7 is blank and skipped; the formula in D8 has no cached value.
    assert [row for row, _ in window["rows"]] == [1, 2, 3, 4, 5, 6, 8]
    assert window["rows"][-1] == [8, [None, "Total", None, None]]


def test_the_verbatim_window_skips_blank_rows_and_is_bounded(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "Report"
    for row in range(4, 40):
        for col in range(1, 30):
            sheet.cell(row=row, column=col, value="x" * 50)
    path = tmp_path / "big.xlsx"
    book.save(path)

    window = WorkbookReader(WORKBOOK_ID, path).describe()[0]["structure"]["window"]

    assert [row for row, _ in window["rows"]] == [1, 4, 5, 6, 7, 8, 9, 10]
    assert len(window["columns"]) == 12
    assert all(len(value) <= 24 for _, values in window["rows"] for value in values if value)


def test_the_sheet_tool_counts_exactly_the_rows_it_would_return(scores_path: Path) -> None:
    # 4 columns; with header_row=2 the data rows are 3..8, six of them: 24 cells.
    exact = ToolRegistry([WorkbookReader(WORKBOOK_ID, scores_path)], max_read_cells=24)
    short = ToolRegistry([WorkbookReader(WORKBOOK_ID, scores_path)], max_read_cells=23)

    fits = _dispatch(exact, "get_sheet_as_dataframe", sheet_name="Scores", header_row=2)
    over = _dispatch(short, "get_sheet_as_dataframe", sheet_name="Scores", header_row=2)
    capped = _dispatch(short, "get_sheet_as_dataframe", sheet_name="Scores", max_rows=2)

    assert fits.state is not ResultState.ERROR, fits.error
    assert over.state is ResultState.ERROR and "24 cells" in (over.error or "")
    assert capped.state is not ResultState.ERROR, capped.error


def test_a_range_may_name_its_sheet(scores_path: Path) -> None:
    # Models write wb.range('Scores!B3:D3') by habit; a real run lost a turn to it.
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, scores_path)])

    named = _dispatch(registry, "inspect_range", range_ref="Scores!B3:D3")
    both = _dispatch(registry, "inspect_range", range_ref="'Scores'!B3:D3", sheet_name="Scores")
    clash = _dispatch(registry, "inspect_range", range_ref="Other!B3", sheet_name="Scores")

    assert named.state is ResultState.SUCCESS, named.error
    assert both.state is ResultState.SUCCESS, both.error
    assert clash.state is ResultState.ERROR and "Other" in (clash.error or "")
