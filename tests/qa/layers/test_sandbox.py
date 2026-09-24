"""Bounded sandbox session behavior; see ``specs/layers/sandbox.md``."""

from collections.abc import Iterator
from pathlib import Path

import pytest
from openpyxl import Workbook

from autotab.qa.layers.sandbox import Sandbox, memory_limit_supported
from autotab.qa.layers.tools import ToolRegistry, WorkbookReader
from autotab.qa.schemas import ResultState

WORKBOOK_ID = "workbook_1"


@pytest.fixture(scope="module")
def workbook_path(tmp_path_factory: pytest.TempPathFactory) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "Revenue"
    sheet["A1"] = "Region"
    sheet["B1"] = "Amount"
    sheet["A2"] = "North"
    sheet["B2"] = 50
    sheet["A3"] = "South"
    sheet["B3"] = 40
    path = tmp_path_factory.mktemp("qa") / "book.xlsx"
    book.save(path)
    return path


@pytest.fixture
def sandbox(workbook_path: Path) -> Iterator[Sandbox]:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, workbook_path, max_range_cells=50)])
    with Sandbox(registry, timeout_seconds=10) as box:
        yield box


def test_names_persist_across_actions_within_one_run(sandbox: Sandbox) -> None:
    first = sandbox.execute("values = [50, 40]", "exec_1")
    second = sandbox.execute("result = sum(values)\nprint(result)", "exec_2")

    assert first.state is ResultState.EMPTY
    assert second.state is ResultState.SUCCESS
    assert second.result == 90
    assert second.stdout.strip() == "90"
    assert "values" in sandbox.bound_names()


def test_import_is_rejected_before_execution(sandbox: Sandbox) -> None:
    sandbox.execute("marker = 1", "exec_1")

    rejected = sandbox.execute("import os\nmarker = 2", "exec_2")
    survivor = sandbox.execute("result = marker", "exec_3")

    assert rejected.state is ResultState.ERROR
    assert "import" in (rejected.error or "").lower()
    assert survivor.result == 1


@pytest.mark.parametrize(
    "code",
    [
        "open('/etc/passwd')",
        "eval('1+1')",
        "__import__('os')",
        "getattr(wb, 'sheets')()",
        "exec('x = 1')",
        "globals()",
    ],
)
def test_forbidden_capabilities_are_absent(sandbox: Sandbox, code: str) -> None:
    result = sandbox.execute(code, "exec_1")

    assert result.state is ResultState.ERROR
    assert result.error


def test_oversized_code_is_refused_without_running(sandbox: Sandbox) -> None:
    result = sandbox.execute("result = 1\n" + "# padding\n" * 5000, "exec_1")

    assert result.state is ResultState.ERROR
    assert "size" in (result.error or "").lower() or "chars" in (result.error or "").lower()


def test_timeout_kills_the_worker_and_the_next_action_starts_clean(workbook_path: Path) -> None:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, workbook_path)])
    with Sandbox(registry, timeout_seconds=1) as box:
        box.execute("keep = 'me'", "exec_1")
        timed_out = box.execute("total = sum(range(10 ** 10))", "exec_2")
        after = box.execute("result = keep", "exec_3")
        alive = box.execute("result = 1 + 1", "exec_4")

    assert timed_out.state is ResultState.ERROR
    assert "timed out" in (timed_out.error or "").lower()
    # The restarted session is usable again, but it does not resume the old namespace.
    assert after.state is ResultState.ERROR
    assert "NameError" in (after.error or "")
    assert alive.result == 2


@pytest.mark.parametrize(
    ("code", "expected"),
    [
        ("while True:\n    x = 1", "while"),
        ("def helper():\n    return 1", "function"),
        ("f = lambda x: x", "lambda"),
        ("class Thing:\n    pass", "class"),
    ],
)
def test_unbounded_constructs_are_refused_statically(
    sandbox: Sandbox, code: str, expected: str
) -> None:
    # These make runtime cost independent of the data, which would leave the
    # wall-clock timeout deciding outcomes differently on different machines.
    result = sandbox.execute(code, "exec_1")

    assert result.state is ResultState.ERROR
    assert expected in (result.error or "").lower()


@pytest.mark.parametrize(
    "code",
    [
        "result = ().__class__.__bases__[0].__subclasses__()",
        "result = wb.__dict__",
        "result = __builtins__",
    ],
)
def test_object_model_access_is_refused(sandbox: Sandbox, code: str) -> None:
    # Restricted builtins alone are not a boundary: dunder access walks from any
    # literal to every loaded class, and from there to os.
    result = sandbox.execute(code, "exec_1")

    assert result.state is ResultState.ERROR
    assert "not available" in (result.error or "")


def test_ordinary_attribute_access_still_works(sandbox: Sandbox) -> None:
    result = sandbox.execute("df = wb.sheet('Revenue')\nresult = int(df['Amount'].sum())", "e1")

    assert result.state is ResultState.SUCCESS
    assert result.result == 90


def test_oversized_syntax_tree_is_refused(workbook_path: Path) -> None:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, workbook_path)])
    with Sandbox(registry, max_nodes=50) as box:
        result = box.execute("\n".join(f"v{i} = {i} + 1" for i in range(40)), "exec_1")

    assert result.state is ResultState.ERROR
    assert "syntax nodes" in (result.error or "")


def test_a_dying_worker_cannot_answer_for_its_replacement(sandbox: Sandbox) -> None:
    # A restart replaces the reply queue while the previous worker's pump thread may
    # still be draining its stdout. If that thread wrote into the live queue, the next
    # action would consume the dead worker's answer.
    retired = sandbox._replies
    process = sandbox._process
    assert process is not None

    sandbox._restart()
    Sandbox._pump(process, retired)

    assert sandbox._replies.empty()
    assert sandbox.execute("result = 'fresh'", "exec_2").result == "fresh"


def test_closing_twice_is_harmless(sandbox: Sandbox) -> None:
    sandbox.close()
    sandbox.close()

    assert sandbox._process is None


def test_facade_reads_go_through_the_tool_registry(sandbox: Sandbox) -> None:
    frame = sandbox.execute(
        "df = wb.sheet('Revenue')\nresult = [list(df.columns), len(df)]", "exec_1"
    )
    cells = sandbox.execute("result = wb.range('B2')[0][0]['value']", "exec_2")

    assert frame.result == [["Region", "Amount"], 2]
    assert cells.result == 50
    assert cells.facade_calls == [
        {"method": "range", "workbook_id": WORKBOOK_ID, "range_ref": "B2", "sheet_name": None}
    ]


def test_facade_enforces_the_same_limits_as_a_tool_call(sandbox: Sandbox) -> None:
    result = sandbox.execute("result = wb.range('A1:Z100')", "exec_1")

    assert result.state is ResultState.ERROR
    assert "50" in (result.error or "")


@pytest.fixture(scope="module")
def ledger_path(tmp_path_factory: pytest.TempPathFactory) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "Ledger"
    sheet["A1"] = "Item"
    sheet["B1"] = "Amount"
    for row, (item, amount) in enumerate([("a", 50), ("b", 40), ("c", None), ("d", 10)], 2):
        sheet[f"A{row}"] = item
        sheet[f"B{row}"] = amount
    sheet["A6"] = "total"
    sheet["B6"] = "=SUM(B2:B5)"  # saved by openpyxl, so it has no cached value
    sheet["D1"] = "Merged"
    sheet.merge_cells("D1:E1")
    book.create_sheet("Other Sheet")["A1"] = 7
    path = tmp_path_factory.mktemp("ledger") / "ledger.xlsx"
    book.save(path)
    return path


@pytest.fixture
def ledger(ledger_path: Path) -> Iterator[Sandbox]:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, ledger_path, max_range_cells=50)])
    with Sandbox(registry, timeout_seconds=10) as box:
        yield box


def _inputs(box: Sandbox, computation_id: str) -> list[tuple[str | None, object]]:
    return [(item.reference, item.value) for item in box.computations[computation_id].inputs]


def test_input_values_are_read_from_the_workbook_not_declared(ledger: Sandbox) -> None:
    result = ledger.execute(
        "result = record_computation('sum', ['Ledger!B2', 'Ledger!B3'], 90, unit='VND')",
        "exec_1",
    )

    assert result.state is ResultState.SUCCESS, result.error
    assert result.computation_ids == [result.result]
    assert _inputs(ledger, str(result.result)) == [("Ledger!B2", 50), ("Ledger!B3", 40)]
    record = ledger.computations[str(result.result)]
    assert record.output == 90
    assert record.metadata["unit"] == "VND"


def test_a_range_expands_to_its_cells_and_records_skipped_blanks(ledger: Sandbox) -> None:
    result = ledger.execute("result = record_computation('max', ['Ledger!B2:B5'], 50)", "e1")

    computation_id = str(result.result)
    assert _inputs(ledger, computation_id) == [
        ("Ledger!B2", 50),
        ("Ledger!B3", 40),
        ("Ledger!B5", 10),
    ]
    assert ledger.computations[computation_id].metadata["skipped_blank"] == ["Ledger!B4"]


def test_a_declared_value_is_refused(ledger: Sandbox) -> None:
    # The model once "proved" a maximum with max([99]) = 99 labelled as a whole column.
    result = ledger.execute(
        "record_computation('max', [{'reference': 'Ledger!B2:B5', 'value': 99}], 99)", "e1"
    )

    assert result.state is ResultState.ERROR
    assert "reads workbook values itself" in (result.error or "")
    assert ledger.computations == {}


def test_constants_alone_prove_nothing(ledger: Sandbox) -> None:
    result = ledger.execute("record_computation('max', [{'constant': 99}], 99)", "e1")

    assert result.state is ResultState.ERROR
    assert "at least one workbook reference" in (result.error or "")


def test_earlier_computations_and_constants_are_explicit_inputs(ledger: Sandbox) -> None:
    ledger.execute("total = record_computation('sum', ['Ledger!B2:B3'], 90)", "e1")
    result = ledger.execute(
        "result = record_computation('ratio', [total, {'constant': 100}], 0.9)", "e2"
    )

    assert result.state is ResultState.SUCCESS, result.error
    assert _inputs(ledger, str(result.result)) == [("calc_1", 90), (None, 100)]


@pytest.mark.parametrize(
    ("reference", "message"),
    [
        ("Ledger!B6", "no cached value"),
        ("Ledger!B4", "blank"),
        ("Nope!A1", "Nope"),
        ("B2", "Sheet!A1"),
        ("calc_9", "calc_9"),
    ],
)
def test_unusable_references_are_refused(ledger: Sandbox, reference: str, message: str) -> None:
    result = ledger.execute(f"record_computation('sum', [{reference!r}], 0)", "e1")

    assert result.state is ResultState.ERROR
    assert message in (result.error or "")


def test_quoted_sheet_names_resolve(ledger: Sandbox) -> None:
    result = ledger.execute("result = record_computation('sum', [\"'Other Sheet'!A1\"], 7)", "e1")

    assert _inputs(ledger, str(result.result)) == [("'Other Sheet'!A1", 7)]


def test_a_merged_range_counts_once(ledger: Sandbox) -> None:
    result = ledger.execute("result = record_computation('count', ['Ledger!D1:E1'], 1)", "e1")

    assert _inputs(ledger, str(result.result)) == [("Ledger!D1", "Merged")]


def test_reads_made_while_recording_are_logged_and_bounded(ledger: Sandbox) -> None:
    logged = ledger.execute("record_computation('sum', ['Ledger!B2:B3'], 90)", "e1")
    oversized = ledger.execute("record_computation('sum', ['Ledger!A1:Z100'], 0)", "e2")

    assert {
        "method": "range",
        "workbook_id": WORKBOOK_ID,
        "range_ref": "B2:B3",
        "sheet_name": "Ledger",
    } in logged.facade_calls
    assert oversized.state is ResultState.ERROR
    assert "50" in (oversized.error or "")


def test_sets_are_ordered_and_non_finite_floats_are_refused(sandbox: Sandbox) -> None:
    ordered = sandbox.execute("result = {3, 1, 2}", "exec_1")
    non_finite = sandbox.execute("result = float('inf')", "exec_2")

    assert ordered.result == [1, 2, 3]
    assert non_finite.state is ResultState.ERROR
    assert "finite" in (non_finite.error or "").lower()


def test_oversized_output_is_truncated_explicitly(workbook_path: Path) -> None:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, workbook_path)])
    with Sandbox(registry, max_output_chars=100) as box:
        result = box.execute("print('x' * 500)", "exec_1")

    assert result.state is ResultState.TRUNCATED
    assert len(result.stdout) <= 100


def test_runtime_error_keeps_the_session_and_hides_the_stack(sandbox: Sandbox) -> None:
    sandbox.execute("kept = 7", "exec_1")
    failed = sandbox.execute("result = 1 / 0", "exec_2")
    after = sandbox.execute("result = kept", "exec_3")

    assert failed.state is ResultState.ERROR
    assert "ZeroDivisionError" in (failed.error or "")
    assert "Traceback" not in (failed.error or "")
    assert after.result == 7


def test_non_ascii_sheet_names_survive_the_worker_protocol(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Doanh thu quý 1"
    sheet["A1"] = "Vùng"
    sheet["A2"] = "Miền Bắc"
    path = tmp_path / "vn.xlsx"
    book.save(path)

    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, path)])
    with Sandbox(registry) as box:
        result = box.execute(
            "df = wb.sheet('Doanh thu quý 1')\nresult = [list(df.columns), df['Vùng'][0]]",
            "exec_1",
        )

    assert result.state is ResultState.SUCCESS
    assert result.result == [["Vùng"], "Miền Bắc"]


def test_worker_starts_where_resource_is_unavailable(monkeypatch: pytest.MonkeyPatch) -> None:
    # `resource` is POSIX-only. A top-level import would stop the worker from
    # starting at all on Windows, which would take the whole sandbox with it.
    import importlib
    import sys

    from autotab.qa.layers import _sandbox_worker

    monkeypatch.setitem(sys.modules, "resource", None)
    try:
        windows_like = importlib.reload(_sandbox_worker)
        assert windows_like._limit_memory(128) is False
    finally:
        monkeypatch.undo()
        importlib.reload(_sandbox_worker)


def test_memory_limit_is_reported_and_enforced_where_the_platform_allows(
    workbook_path: Path,
) -> None:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, workbook_path)])
    with Sandbox(registry, memory_limit_mb=128, timeout_seconds=20) as box:
        assert box.memory_limit_enforced is memory_limit_supported()
        if not box.memory_limit_enforced:
            pytest.skip("this platform refuses to lower the address-space limit")
        result = box.execute("result = len(bytearray(400 * 1024 * 1024))", "exec_1")

    assert result.state is ResultState.ERROR
    assert "memory" in (result.error or "").lower()


def test_a_large_frame_result_is_previewed_and_stays_bound(workbook_path: Path) -> None:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, workbook_path)])
    with Sandbox(registry, max_inline_cells=20) as box:
        big = box.execute(
            "frame = pd.DataFrame({'a': list(range(100)), 'b': list(range(100))})\n"
            "result = frame",
            "exec_1",
        )
        later = box.execute("result = int(frame['a'].sum())", "exec_2")

    # Showing less is not a failed computation: the execution state stays SUCCESS.
    assert big.state is ResultState.SUCCESS
    assert isinstance(big.result, dict)
    assert big.result["preview"] is True
    assert big.result["shape"] == [100, 2]
    assert big.result["head"][0] == [0, 0] and big.result["tail"][-1] == [99, 99]
    assert later.result == 4950


def test_a_long_list_result_is_previewed(workbook_path: Path) -> None:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, workbook_path)])
    with Sandbox(registry, max_inline_cells=20) as box:
        result = box.execute("result = list(range(1000))", "exec_1")

    assert isinstance(result.result, dict)
    assert result.result["length"] == 1000
    assert result.result["head"] == [0, 1, 2, 3, 4]
    assert result.result["tail"] == [995, 996, 997, 998, 999]


def test_long_stdout_keeps_both_ends(workbook_path: Path) -> None:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, workbook_path)])
    with Sandbox(registry, max_output_chars=200) as box:
        result = box.execute("for i in range(500):\n    print('line', i)", "exec_1")

    assert result.state is ResultState.TRUNCATED
    assert "line 0" in result.stdout and "line 499" in result.stdout
    assert "characters omitted" in result.stdout
    assert len(result.stdout) <= 200


def test_a_numeric_operation_needs_a_numeric_output(ledger: Sandbox) -> None:
    # A real run recorded a "max" whose output was ['Nga Anh Ngo']: nothing could
    # replay it, and a name claim got linked to it.
    result = ledger.execute("record_computation('max', ['Ledger!B2:B3'], ['a'])", "e1")

    assert result.state is ResultState.ERROR
    assert "number" in (result.error or "") and "citations" in (result.error or "")


def test_an_unknown_operation_is_refused_with_the_list(ledger: Sandbox) -> None:
    result = ledger.execute("record_computation('argmax', ['Ledger!B2:B3'], 'B2')", "e1")

    assert result.state is ResultState.ERROR
    assert "argmax" in (result.error or "") and "sum" in (result.error or "")


@pytest.mark.parametrize(
    ("code", "expected"),
    [
        ("import pandas as pd\nresult = int(pd.Series([1, 2]).sum())", 3),
        ("from datetime import date\nresult = str(date(2020, 1, 2))", "2020-01-02"),
        ("import math, json\nresult = json.dumps(math.floor(2.5))", "2"),
    ],
)
def test_the_already_bound_modules_can_be_imported(
    sandbox: Sandbox, code: str, expected: object
) -> None:
    # Models write `import pandas as pd` by habit; refusing it cost a turn in most runs.
    result = sandbox.execute(code, "exec_1")

    assert result.state is ResultState.SUCCESS, result.error
    assert result.result == expected


@pytest.mark.parametrize(
    "code",
    [
        "import numpy",
        "import os.path",
        "from pandas import read_csv",
        "from pandas.io import common",
    ],
)
def test_other_imports_stay_refused(sandbox: Sandbox, code: str) -> None:
    result = sandbox.execute(code, "exec_1")

    assert result.state is ResultState.ERROR
    assert "not available" in (result.error or "")


@pytest.mark.parametrize(
    "code",
    [
        "df = pd.read_csv('/etc/hosts')",
        "pd.DataFrame({'a': [1]}).to_csv('out.csv')",
        "pd.DataFrame({'a': [1]}).to_excel('out.xlsx')",
        "w = pd.ExcelWriter('out.xlsx')",
        "c = pd.io.common",
    ],
)
def test_pandas_file_access_is_refused(sandbox: Sandbox, code: str) -> None:
    # pandas can read and write arbitrary paths; the workbook is read through wb only.
    result = sandbox.execute(code, "exec_1")

    assert result.state is ResultState.ERROR
    assert "file access" in (result.error or "")


def test_the_worker_runs_in_a_scratch_directory(workbook_path: Path) -> None:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, workbook_path)])
    with Sandbox(registry) as box:
        workdir = box.workdir
        assert workdir is not None and workdir.is_dir()
        assert Path.cwd() not in [workdir, *workdir.parents]

    assert not workdir.exists()


def test_missing_values_inside_frames_come_back_as_null(sandbox: Sandbox) -> None:
    # Blank cells become NaN in pandas; a real run lost a turn when a frame holding
    # them was refused as a "non-finite float" although the code had worked.
    frame = sandbox.execute("result = pd.DataFrame({'a': [1, None], 'b': ['x', None]})", "e1")
    series = sandbox.execute("result = pd.Series([2.5, None])", "e2")

    assert frame.state is ResultState.SUCCESS, frame.error
    assert frame.result == {"columns": ["a", "b"], "rows": [[1.0, "x"], [None, None]]}
    assert series.result == [2.5, None]


def test_a_large_frame_preview_also_nulls_missing_values(workbook_path: Path) -> None:
    registry = ToolRegistry([WorkbookReader(WORKBOOK_ID, workbook_path)])
    with Sandbox(registry, max_inline_cells=4) as box:
        result = box.execute("result = pd.DataFrame({'a': [None] * 10})", "e1")

    assert result.state is ResultState.SUCCESS, result.error
    assert isinstance(result.result, dict) and result.result["head"][0] == [None]


def test_the_facade_names_the_sheet_argument_like_every_tool(sandbox: Sandbox) -> None:
    # wb.sheet alone used `name=` while wb.range, wb.search, and every tool use
    # `sheet_name=`; a real run lost two turns to that mismatch.
    keyword = sandbox.execute("result = len(wb.sheet(sheet_name='Revenue'))", "e1")
    positional = sandbox.execute("result = len(wb.sheet('Revenue'))", "e2")

    assert keyword.state is ResultState.SUCCESS, keyword.error
    assert keyword.result == positional.result == 2
